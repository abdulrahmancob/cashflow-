"""Analytics API role gates (no live DB)."""

from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path

from fastapi import FastAPI
from fastapi.testclient import TestClient

from cashflow_db.repository.work_analytics import TEAM_SS, resolve_team
from cashflow_ops.security import AuthUser, get_current_user
from cashflow_ops.work_analytics_api import VIEW_ROLES, router

USER_ID = "11111111-1111-1111-1111-111111111111"
OTHER_ID = "22222222-2222-2222-2222-222222222222"


def _user(*roles: str) -> AuthUser:
    return AuthUser(
        user_id=USER_ID,
        username="jane@example.com",
        display_name="Jane",
        roles=list(roles),
    )


class _Conn:
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def _app(
    user: AuthUser,
    monkeypatch,
    *,
    summary=None,
    detail=None,
    heartbeat=None,
    breakdown=None,
):
    @contextmanager
    def _connection():
        yield _Conn()

    captured: dict = {}

    def fake_summary(conn, roles, **k):
        captured["team"] = k.get("team")
        captured["roles"] = roles
        resolve_team(roles, k.get("team"))
        return summary or {
            "team": TEAM_SS,
            "kpis": {"people": 0},
            "people": [],
            "charts": {"hours": [], "outcomes": []},
        }

    def fake_breakdown(conn, roles, **k):
        captured["breakdown"] = k
        resolve_team(roles, TEAM_SS)
        grain = (k.get("grain") or "month").strip().lower()
        if grain not in {"month", "week", "day"}:
            raise ValueError("grain must be month, week, or day")
        return breakdown or {
            "grain": grain,
            "year": k.get("year") or 2026,
            "month": None,
            "rows": [{"period": f"Jan {k.get('year') or 2026}"} for _ in range(12)],
            "totals": {"period": "Total", "claims": 0},
            "member": None,
            "people": [],
        }

    monkeypatch.setattr("cashflow_db.repository.connection", _connection)
    monkeypatch.setattr(
        "cashflow_db.repository.work_analytics.record_heartbeat",
        lambda *a, **k: heartbeat or {"ok": True, "counted": True, "action": "open"},
    )
    monkeypatch.setattr("cashflow_db.repository.work_analytics.team_summary", fake_summary)
    monkeypatch.setattr(
        "cashflow_db.repository.work_analytics.user_detail",
        lambda *a, **k: detail,
    )
    monkeypatch.setattr("cashflow_db.repository.work_analytics.ss_breakdown", fake_breakdown)

    app = FastAPI()
    app.include_router(router, prefix="/api")
    app.dependency_overrides[get_current_user] = lambda: user
    app.state.captured = captured
    return app


def test_view_roles_are_admin_and_leads():
    assert "super_admin" in VIEW_ROLES
    assert "ops_admin" in VIEW_ROLES
    assert "second_submission_lead" in VIEW_ROLES
    assert "sub_admin" in VIEW_ROLES
    assert "analytics_viewer" in VIEW_ROLES
    assert "posting_team" not in VIEW_ROLES
    assert "finance" not in VIEW_ROLES
    assert "second_submission" not in VIEW_ROLES


def test_my_today_is_personal_and_rejects_unknown_area(monkeypatch):
    monkeypatch.setattr(
        "cashflow_db.repository.work_analytics.completed_today_for_user",
        lambda conn, user_id, area, **k: 7 if area == "eligibility" else 0,
    )
    client = TestClient(_app(_user("posting_team"), monkeypatch))
    res = client.get("/api/analytics/my-today?area=eligibility")
    assert res.status_code == 200
    assert res.json()["completed_today"] == 7

    def _reject(conn, user_id, area, **k):
        raise ValueError("area must be eligibility, collection, or submission")

    monkeypatch.setattr(
        "cashflow_db.repository.work_analytics.completed_today_for_user",
        _reject,
    )
    res = client.get("/api/analytics/my-today?area=finance")
    assert res.status_code == 400


def test_heartbeat_allows_any_authenticated_user(monkeypatch):
    client = TestClient(_app(_user("posting_team"), monkeypatch))
    res = client.post("/api/analytics/heartbeat", json={"page_path": "/eligibility"})
    assert res.status_code == 200
    assert res.json()["ok"] is True


def test_heartbeat_forwards_presence(monkeypatch):
    seen: dict = {}

    def fake(conn, user_id, **kwargs):
        seen.update(kwargs)
        return {"ok": True, "counted": True, "action": "extend", "add_seconds": 0}

    client = TestClient(_app(_user("collector"), monkeypatch))
    monkeypatch.setattr("cashflow_db.repository.work_analytics.record_heartbeat", fake)
    res = client.post("/api/analytics/heartbeat", json={"presence": True, "idle": False})
    assert res.status_code == 200
    assert seen["presence"] is True
    assert seen["idle"] is False
    res = client.post(
        "/api/analytics/heartbeat",
        json={"presence": True, "desk_permission": "denied"},
    )
    assert res.status_code == 200
    assert seen["desk_permission"] == "denied"


def test_heartbeat_forwards_closed(monkeypatch):
    seen: dict = {}

    def fake(conn, user_id, **kwargs):
        seen.update(kwargs)
        return {"ok": True, "counted": True, "action": "close", "add_seconds": 0}

    client = TestClient(_app(_user("collector"), monkeypatch))
    monkeypatch.setattr("cashflow_db.repository.work_analytics.record_heartbeat", fake)
    res = client.post("/api/analytics/heartbeat", json={"closed": True})
    assert res.status_code == 200
    assert seen["closed"] is True
    assert seen["presence"] is False


def test_team_forbidden_for_posting_finance_ss(monkeypatch):
    for role in ("posting_team", "finance", "second_submission", "collector"):
        client = TestClient(_app(_user(role), monkeypatch))
        res = client.get("/api/analytics/team")
        assert res.status_code == 403, role


def test_team_ok_for_ops_admin_and_ss_lead(monkeypatch):
    summary = {"kpis": {"people": 1, "ss_claims": 2}, "people": [], "charts": {"hours": []}}
    for role in ("ops_admin", "second_submission_lead", "super_admin", "sub_admin", "analytics_viewer"):
        client = TestClient(_app(_user(role), monkeypatch, summary=summary))
        res = client.get("/api/analytics/team?preset=today&team=second_submission")
        assert res.status_code == 200, role
        assert res.json()["kpis"]["ss_claims"] == 2


def test_team_unknown_rejected(monkeypatch):
    client = TestClient(_app(_user("ops_admin"), monkeypatch))
    res = client.get("/api/analytics/team?team=tracker")
    assert res.status_code == 400


def test_ss_lead_cannot_request_other_known_team(monkeypatch):
    import cashflow_db.repository.work_analytics as wa

    monkeypatch.setattr(wa, "KNOWN_TEAMS", (TEAM_SS, "eligibility"))
    client = TestClient(_app(_user("second_submission_lead"), monkeypatch))
    res = client.get("/api/analytics/team?team=eligibility")
    assert res.status_code == 403


def test_ss_breakdown_month_ok(monkeypatch):
    client = TestClient(_app(_user("ops_admin"), monkeypatch))
    res = client.get("/api/analytics/ss/breakdown?grain=month&year=2026")
    assert res.status_code == 200
    body = res.json()
    assert body["grain"] == "month"
    assert len(body["rows"]) == 12
    assert body["totals"]["period"] == "Total"


def test_collection_root_causes_ok(monkeypatch):
    monkeypatch.setattr(
        "cashflow_db.repository.work_analytics.collection_root_cause_breakdown",
        lambda conn, roles, **k: {
            "year": k.get("year") or 2026,
            "labels": ["Auth delay"],
            "rows": [{"period": "Jan"} for _ in range(12)],
        },
    )
    client = TestClient(_app(_user("ops_admin"), monkeypatch))
    res = client.get("/api/analytics/collection/root-causes?year=2026")
    assert res.status_code == 200
    body = res.json()
    assert body["year"] == 2026
    assert len(body["rows"]) == 12
    assert body["labels"] == ["Auth delay"]


def test_my_assignments_for_collector(monkeypatch):
    monkeypatch.setattr(
        "cashflow_db.repository.work_analytics.my_assignment_progress",
        lambda conn, user_id: {"assigned": 20, "finished": 8},
    )
    client = TestClient(_app(_user("collector"), monkeypatch))
    res = client.get("/api/analytics/my-assignments")
    assert res.status_code == 200
    assert res.json() == {"assigned": 20, "finished": 8}


def test_collection_dead_root_causes_ok(monkeypatch):
    monkeypatch.setattr(
        "cashflow_db.repository.work_analytics.collection_dead_root_causes",
        lambda conn, roles: {
            "total": 4,
            "rows": [{"label": "Auth delay", "count": 3, "percent": 75.0}],
        },
    )
    client = TestClient(_app(_user("ops_admin"), monkeypatch))
    res = client.get("/api/analytics/collection/dead-root-causes")
    assert res.status_code == 200
    body = res.json()
    assert body["total"] == 4
    assert body["rows"][0]["label"] == "Auth delay"
    assert float(body["rows"][0]["percent"]) == 75.0
    page = (
        Path(__file__).resolve().parents[2] / "rcm_portal" / "src" / "pages" / "TeamAnalytics.tsx"
    ).read_text(encoding="utf-8")
    assert "Dead by root cause" in page
    assert 'layout="vertical"' in page


def test_ss_lead_cannot_open_collection_root_causes(monkeypatch):
    def _reject(conn, roles, **k):
        resolve_team(roles, "collection")

    monkeypatch.setattr(
        "cashflow_db.repository.work_analytics.collection_root_cause_breakdown",
        _reject,
    )
    client = TestClient(_app(_user("second_submission_lead"), monkeypatch))
    res = client.get("/api/analytics/collection/root-causes")
    assert res.status_code == 403


def test_ss_breakdown_bad_grain(monkeypatch):
    client = TestClient(_app(_user("ops_admin"), monkeypatch))
    res = client.get("/api/analytics/ss/breakdown?grain=hour")
    assert res.status_code == 400


def test_export_streams_workbook_for_view_roles_only(monkeypatch):
    from io import BytesIO

    from openpyxl import load_workbook

    monkeypatch.setattr(
        "cashflow_db.repository.work_analytics.collection_root_cause_breakdown",
        lambda conn, roles, **k: {
            "year": 2026,
            "labels": ["Auth delay"],
            "rows": [],
            "people": [],
        },
    )
    monkeypatch.setattr(
        "cashflow_db.repository.work_analytics.collection_dead_root_causes",
        lambda conn, roles: {"total": 0, "rows": []},
    )
    page = (
        Path(__file__).resolve().parents[2] / "rcm_portal" / "src" / "pages" / "TeamAnalytics.tsx"
    ).read_text(encoding="utf-8")
    assert "Download sheet" in page
    analytics_ts = (
        Path(__file__).resolve().parents[2] / "rcm_portal" / "src" / "api" / "analytics.ts"
    ).read_text(encoding="utf-8")
    assert "ss-team-analytics.xlsx" in analytics_ts
    assert "p.set('team', opts.team)" in analytics_ts
    assert "exportCsv" not in page

    for role in ("posting_team", "finance", "second_submission", "collector"):
        client = TestClient(_app(_user(role), monkeypatch))
        res = client.get("/api/analytics/export")
        assert res.status_code == 403, role

    client = TestClient(_app(_user("ops_admin"), monkeypatch))
    res = client.get("/api/analytics/export?team=second_submission&preset=month&year=2026&grain=month")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    assert "ss-team-analytics.xlsx" in res.headers["content-disposition"]
    book = load_workbook(BytesIO(res.content))
    assert book.sheetnames == [
        "SS Summary",
        "SS Hours",
        "SS Outcomes",
        "SS People",
        "SS Claim analysis",
    ]

    coll = client.get("/api/analytics/export?team=collection")
    assert coll.status_code == 200
    assert "collection-team-analytics.xlsx" in coll.headers["content-disposition"]
    coll_book = load_workbook(BytesIO(coll.content))
    assert "Collection Dead" in coll_book.sheetnames
    assert "Collection Root causes" in coll_book.sheetnames
    assert "SS Summary" not in coll_book.sheetnames

    lead = TestClient(_app(_user("second_submission_lead"), monkeypatch))
    lead_res = lead.get("/api/analytics/export?team=second_submission")
    assert lead_res.status_code == 200
    lead_book = load_workbook(BytesIO(lead_res.content))
    assert "SS People" in lead_book.sheetnames
    assert "Collection Summary" not in lead_book.sheetnames
    assert lead.get("/api/analytics/export?team=collection").status_code == 403

    bad = client.get("/api/analytics/export?user_id=not-a-uuid")
    assert bad.status_code == 400


def test_team_user_404_when_out_of_scope(monkeypatch):
    client = TestClient(_app(_user("ops_admin"), monkeypatch, detail=None))
    res = client.get(f"/api/analytics/team/{OTHER_ID}")
    assert res.status_code == 404
