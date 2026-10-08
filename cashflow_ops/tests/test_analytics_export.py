"""Team Analytics workbook matches the dashboard sections."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook

from cashflow_ops.analytics_export import export_tables, format_hours, format_money


def _sheet(tables, title):
    match = next(item for item in tables if item[0] == title)
    return match[1], match[2]


def test_formats_match_the_dashboard():
    assert format_hours(0) == "0h"
    assert format_hours(30) == "<1m"
    assert format_hours(3660) == "1h 1m"
    assert format_hours(7200) == "2h"
    assert format_money(1234) == "$1,234"
    assert format_money(-20) == "-$20"


def test_workbook_includes_every_dashboard_section():
    summaries = {
        "second_submission": {
            "period": {"start": "2026-10-01", "end": "2026-10-08"},
            "kpis": {
                "ss_claims_today": 1,
                "ss_claims_week": 2,
                "ss_claims_month": 3,
                "ss_money": 1234,
                "ss_paid": 4,
                "ss_denied": 5,
                "ss_timely_filing": 6,
            },
            "people": [
                {
                    "display_name": "Ada",
                    "roles": ["second_submission"],
                    "seconds_today": 3660,
                    "seconds_week": 0,
                    "seconds_month": 7200,
                    "ss_claims_today": 1,
                    "ss_claims_week": 2,
                    "ss_claims_month": 3,
                    "ss_paid": 4,
                    "ss_denied": 5,
                    "ss_timely_filing": 6,
                }
            ],
            "charts": {
                "hours": [{"name": "Ada", "hours": 2.0}],
                "outcomes": [{"name": "Ada", "paid": 4, "denied": 5, "timely_filing": 6}],
            },
        },
        "eligibility": {
            "period": {"start": "2026-10-01", "end": "2026-10-08"},
            "kpis": {"elig_touched_today": 8, "elig_money": 90},
            "people": [],
            "charts": {"hours": [], "outcomes": []},
        },
        "collection": {
            "period": {"start": "2026-10-01", "end": "2026-10-08"},
            "kpis": {"coll_touched_today": 2, "coll_money_today": 40},
            "collection_statuses": ["Paid"],
            "people": [
                {
                    "display_name": "Bea",
                    "roles": ["collector"],
                    "seconds_today": 60,
                    "seconds_month": 0,
                    "coll_touched_today": 2,
                    "coll_touched_month": 4,
                    "coll_assigned": 5,
                    "coll_finished": 1,
                    "coll_status_today": {"Paid": 3},
                    "coll_status_month": {"Paid": 7},
                    "coll_recovered_today": 1,
                    "coll_recovered_month": 2,
                    "coll_money_today": 40,
                    "coll_money_month": 80,
                }
            ],
            "charts": {"hours": [], "outcomes": []},
        },
        "submission": {
            "period": {"start": "2026-10-01", "end": "2026-10-08"},
            "kpis": {"cpt_touched_today": 1, "icd_touched_today": 2, "cpt_resolved_month": 4},
            "people": [],
            "charts": {"hours": [], "outcomes": []},
        },
    }
    breakdown = {
        "grain": "month",
        "year": 2026,
        "member": None,
        "rows": [
            {
                "period": "Jan 2026",
                "payment": 50,
                "claims": 2,
                "total_submitted_claims": 3,
                "submitted": 1,
                "paid": 1,
                "pending": 0,
                "denied": 0,
                "corrected": 0,
                "timely_filing": 0,
            }
        ],
        "totals": {
            "period": "Total",
            "payment": 50,
            "claims": 2,
            "total_submitted_claims": 3,
            "submitted": 1,
            "paid": 1,
            "pending": 0,
            "denied": 0,
            "corrected": 0,
            "timely_filing": 0,
        },
    }
    causes = {
        "year": 2026,
        "labels": ["Auth delay"],
        "rows": [
            {
                "period": "Jan 2026",
                "counts": {"Auth delay": 2},
                "top": "Auth delay",
                "top_count": 2,
            }
        ],
        "people": [
            {
                "display_name": "Bea",
                "counts": {"Auth delay": 2},
                "top": "Auth delay",
                "top_count": 2,
            }
        ],
    }
    dead = {"total": 4, "rows": [{"label": "Auth delay", "count": 3, "percent": 75.0}]}

    tables = export_tables(
        summaries, breakdown=breakdown, causes=causes, dead=dead
    )
    names = [title for title, _headers, _rows in tables]
    assert names == [
        "SS Summary",
        "SS Hours",
        "SS Outcomes",
        "SS People",
        "SS Claim analysis",
        "Eligibility Summary",
        "Eligibility Hours",
        "Eligibility People",
        "Collection Summary",
        "Collection Hours",
        "Collection People",
        "Collection Root causes",
        "Collection Root people",
        "Collection Dead",
        "Submission Summary",
        "Submission Hours",
        "Submission People",
    ]
    for title, _headers, _rows in tables:
        assert len(title) <= 31

    _headers, summary = _sheet(tables, "SS Summary")
    cards = {row[0]: row[1] for row in summary}
    assert cards["Title"] == "Second Submission"
    assert cards["Claims today"] == 1
    assert cards["Payment"] == "$1,234"
    assert cards["Timely Filing"] == 6
    assert "Submitter" in cards["Description"]
    assert cards["Claim analysis"].startswith("All team")

    _headers, people = _sheet(tables, "SS People")
    assert people[0][0] == "Ada"
    assert people[0][2] == "1h 1m"
    assert "Payment" not in _headers
    assert "Pending" not in _headers

    _headers, claims = _sheet(tables, "SS Claim analysis")
    assert claims[0][0] == "Jan 2026"
    assert claims[0][1] == "$50"
    assert claims[-1][0] == "Total"

    _headers, coll = _sheet(tables, "Collection People")
    assert coll[0][0] == "Bea"
    assert "Paid today" in _headers
    assert "Hours week" not in _headers
    assert coll[0][_headers.index("Paid today")] == 3
    assert coll[0][_headers.index("Money today")] == "$40"

    _headers, causes_rows = _sheet(tables, "Collection Root causes")
    assert causes_rows[0] == ["Jan 2026", "Auth delay (2)", 2]

    _headers, dead_rows = _sheet(tables, "Collection Dead")
    assert dead_rows[0] == ["Auth delay", 3, "75.0%", "3 (75.0%)"]

    _headers, elig = _sheet(tables, "Eligibility Summary")
    assert {row[0]: row[1] for row in elig}["Touched today"] == 8

    workbook = Workbook(write_only=True)
    for title, headers, rows in tables:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    loaded = load_workbook(BytesIO(buffer.getvalue()))
    assert loaded.sheetnames == names
    assert loaded["SS People"]["A2"].value == "Ada"
