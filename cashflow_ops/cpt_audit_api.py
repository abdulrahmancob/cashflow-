"""CPT audit work-queue and finance insight APIs."""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from cashflow_ops.heavy import heavy_guard
from cashflow_ops.security import (
    ROLE_FINANCE,
    ROLE_MEDICAL_AUDIT,
    ROLE_OPS_ADMIN,
    ROLE_SUB_ADMIN,
    ROLE_SUPER,
    AuthUser,
    require_roles,
)

router = APIRouter(prefix="/cpt-audit", tags=["cpt-audit"])

QUEUE_VIEW = (ROLE_MEDICAL_AUDIT, ROLE_SUPER, ROLE_SUB_ADMIN, ROLE_OPS_ADMIN)
QUEUE_EDIT = (ROLE_MEDICAL_AUDIT, ROLE_SUPER, ROLE_SUB_ADMIN, ROLE_OPS_ADMIN)
INSIGHTS_VIEW = (ROLE_FINANCE, ROLE_SUPER, ROLE_SUB_ADMIN)
INSIGHTS_PUBLISH = (ROLE_SUPER, ROLE_SUB_ADMIN)

POSTING_MAX_DAYS = 30
SUPER_MAX_DAYS = 365


def _ser(obj: Any) -> Any:
    if isinstance(obj, dict):
        return {k: _ser(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_ser(v) for v in obj]
    if isinstance(obj, (datetime, date)):
        return obj.isoformat()
    if isinstance(obj, UUID):
        return str(obj)
    if hasattr(obj, "hex"):
        return str(obj)
    return obj


class WorkPatch(BaseModel):
    workflow_status: str | None = None
    resolution_code: str | None = None
    resolution_note: str | None = Field(default=None, max_length=4000)
    ignore_reason: str | None = None
    assigned_to: str | None = None


class PublishBody(BaseModel):
    published: bool = True


def _days_for(user: AuthUser, days: int | None) -> int:
    requested = int(days or POSTING_MAX_DAYS)
    cap = SUPER_MAX_DAYS if user.is_elevated_admin else POSTING_MAX_DAYS
    return max(1, min(requested, cap))


@router.get("/meta")
def meta(
    days: int | None = None,
    user: AuthUser = Depends(require_roles(*QUEUE_VIEW)),
) -> dict[str, Any]:
    from cashflow_db.repository import connection, cpt_audit

    d = _days_for(user, days)
    with connection() as conn:
        return _ser({"filters": cpt_audit.filter_options(conn, days=d), "days": d})


@router.get("/items")
def list_items(
    q: str | None = None,
    bucket: list[str] | None = Query(None),
    rule_code: list[str] | None = Query(None),
    insurance: list[str] | None = Query(None),
    facility: list[str] | None = Query(None),
    cluster: list[str] | None = Query(None),
    status: list[str] | None = Query(None),
    claim_status: list[str] | None = Query(None),
    domain: list[str] | None = Query(None),
    signed_on: date | None = None,
    days: int | None = None,
    page: int = 1,
    page_size: int = 50,
    user: AuthUser = Depends(require_roles(*QUEUE_VIEW)),
) -> dict[str, Any]:
    from cashflow_db.repository import connection, cpt_audit

    d = _days_for(user, days)
    with connection() as conn:
        data = cpt_audit.list_queue(
            conn,
            q=q,
            bucket=bucket,
            rule_code=rule_code,
            insurance=insurance,
            facility=facility,
            cluster=cluster,
            workflow_status=status,
            claim_status=claim_status,
            audit_domain=domain,
            signed_on=signed_on,
            days=d,
            page=page,
            page_size=page_size,
        )
    return _ser(data)


@router.get("/items/export", dependencies=[Depends(heavy_guard)])
def export_items(
    q: str | None = None,
    bucket: list[str] | None = Query(None),
    rule_code: list[str] | None = Query(None),
    insurance: list[str] | None = Query(None),
    facility: list[str] | None = Query(None),
    cluster: list[str] | None = Query(None),
    status: list[str] | None = Query(None),
    claim_status: list[str] | None = Query(None),
    domain: list[str] | None = Query(None),
    signed_on: date | None = None,
    days: int | None = None,
    user: AuthUser = Depends(require_roles(*QUEUE_VIEW)),
) -> StreamingResponse:
    from cashflow_db.repository import connection, cpt_audit

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="openpyxl required") from exc

    d = _days_for(user, days)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Audit Queue")
    headers = cpt_audit.sheet_export_headers()
    ws.append(headers)
    row_count = 1
    page = 1
    page_size = 200
    with connection() as conn:
        while True:
            data = cpt_audit.list_queue(
                conn,
                q=q,
                bucket=bucket,
                rule_code=rule_code,
                insurance=insurance,
                facility=facility,
                cluster=cluster,
                workflow_status=status,
                claim_status=claim_status,
                audit_domain=domain,
                signed_on=signed_on,
                days=d,
                page=page,
                page_size=page_size,
            )
            items = data.get("items") or []
            for row in items:
                ws.append(cpt_audit.sheet_export_row(row))
                row_count += 1
            total = int(data.get("total") or 0)
            if not items or page * page_size >= total:
                break
            page += 1
    if row_count > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_count}"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=audit_queue.xlsx"},
    )


@router.get("/items/{work_item_id}")
def get_item(
    work_item_id: str,
    _: AuthUser = Depends(require_roles(*QUEUE_VIEW)),
) -> dict[str, Any]:
    from cashflow_db.repository import connection, cpt_audit

    with connection() as conn:
        item = cpt_audit.get_item(conn, work_item_id)
        if not item:
            raise HTTPException(status_code=404, detail="Not found")
        history = cpt_audit.list_item_history(conn, work_item_id)
    return _ser({"item": item, "history": history})


@router.patch("/items/{work_item_id}")
def patch_item(
    work_item_id: str,
    body: WorkPatch,
    user: AuthUser = Depends(require_roles(*QUEUE_EDIT)),
) -> dict[str, Any]:
    from cashflow_db.repository import connection, cpt_audit

    payload = body.model_dump(exclude_unset=True)
    try:
        with connection() as conn:
            before = cpt_audit.get_item(conn, work_item_id)
            item = cpt_audit.patch_work_item(
                conn, work_item_id, payload, actor_user_id=user.user_id
            )
            from cashflow_ops.activity_api import log_write
            from cashflow_db.repository import portal_activity

            row = item or before or {}
            name = str(row.get("patient_name") or "Visit").strip() or "Visit"
            dos = portal_activity.pretty_value(row.get("dos")) if row.get("dos") else ""
            rule = str(row.get("rule_code") or "").strip()
            parts = [name]
            if dos and dos != "empty":
                parts.append(dos)
            if rule:
                parts.append(rule)
            parts.append("CPT audit")
            log_write(
                conn,
                user,
                action="updated",
                area="cpt_audit",
                entity_type="cpt_audit_item",
                entity_id=work_item_id,
                entity_label=" · ".join(parts),
                before=before,
                after=item,
                keys=list(payload.keys()),
            )
    except KeyError as exc:
        raise HTTPException(status_code=404, detail="Not found") from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return _ser({"item": item})


@router.get("/insights")
def insights(
    _: AuthUser = Depends(require_roles(*INSIGHTS_VIEW)),
) -> dict[str, Any]:
    from cashflow_db.repository import connection, cpt_audit

    with connection() as conn:
        data = cpt_audit.insight_totals(conn)
    return _ser(data)


@router.post("/insights/publish")
def publish_insights(
    body: PublishBody,
    user: AuthUser = Depends(require_roles(*INSIGHTS_PUBLISH)),
) -> dict[str, Any]:
    from cashflow_db.repository import connection, cpt_audit

    with connection() as conn:
        row = cpt_audit.set_insights_published(
            conn, body.published, actor_user_id=user.user_id
        )
        from cashflow_ops.activity_api import log_write

        log_write(
            conn,
            user,
            action="published",
            area="cpt_audit",
            entity_type="insights",
            entity_id="cpt-insights",
            entity_label="Business Insights",
            after={"published": body.published},
            fallback="Published insights" if body.published else "Unpublished insights",
        )
    return _ser({"settings": row})
