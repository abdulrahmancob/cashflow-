"""Checks & Deposits workbook parse (no Postgres)."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from cashflow_db.loaders.checks_deposits_xlsx import (
    export_checks_deposits_workbook,
    parse_checks_deposits_workbook,
    parse_sheet_money,
)


def _workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PT OF THE CITY"
    ws.append(
        [
            None,
            "Source(insurance)",
            "Amount",
            "Check number ",
            "Deposit Date",
            "Link ",
            "Source(insurance)",
            "Amount(Visa)",
        ]
    )
    ws.append(
        [
            date(2026, 1, 2),
            "CSS LLC",
            827.92,
            69017024.0,
            date(2026, 1, 5),
            "BH-CSS.pdf",
            None,
            None,
        ]
    )
    ws.append(
        [
            date(2026, 1, 3),
            "GEHA",
            "1700,49",
            "FC1",
            date(2026, 1, 6),
            "geha.pdf",
            "ignored",
            99,
        ]
    )
    ws.append(
        [
            date(2026, 2, 1),
            "Bad",
            "103.003.64",
            "BAD1",
            date(2026, 2, 2),
            None,
            None,
            None,
        ]
    )
    ws.append(
        [
            date(2025, 6, 1),
            "Old",
            50,
            "OLD1",
            date(2025, 6, 2),
            None,
            None,
            None,
        ]
    )
    ws.append(
        [
            date(2026, 3, 9),
            "Dup",
            90,
            "117130664J",
            date(2026, 3, 9),
            "a.pdf",
            None,
            None,
        ]
    )
    ws.append(
        [
            date(2026, 3, 9),
            "Dup",
            90,
            "117130664J",
            date(2026, 3, 9),
            "b.pdf",
            None,
            None,
        ]
    )
    ws.append(
        [
            date(2026, 4, 1),
            None,
            None,
            None,
            date(2026, 4, 2),
            None,
            "Visa Payer",
            12,
        ]
    )
    ws.append(
        [
            "03/06/0206",
            "Typo",
            10,
            "1170019231",
            date(2026, 3, 9),
            "typo.pdf",
            None,
            None,
        ]
    )
    ws.append(
        [
            date(2026, 5, 1),
            "Blank",
            " ",
            "BLANK1",
            date(2026, 5, 2),
            None,
            None,
            None,
        ]
    )
    wb.save(path)


def test_parse_sheet_money_comma_decimal_and_ambiguous():
    amount, err = parse_sheet_money("1700,49")
    assert err is None
    assert amount == Decimal("1700.49")
    amount, err = parse_sheet_money("103.003.64")
    assert amount is None
    assert err and "Ambiguous" in err
    assert parse_sheet_money(" ") == (None, None)
    assert parse_sheet_money(827.92)[0] == Decimal("827.92")


def test_parse_workbook_gates_2026_and_ignores_visa(tmp_path: Path):
    path = tmp_path / "checks.xlsx"
    _workbook(path)
    parsed = parse_checks_deposits_workbook(path)
    by_key = {row.sheet_key: row for row in parsed.rows}

    assert "69017024|2026-01-05|827.92|0" in by_key
    assert by_key["69017024|2026-01-05|827.92|0"].check_number == "69017024"
    assert by_key["69017024|2026-01-05|827.92|0"].payer == "CSS LLC"
    assert by_key["69017024|2026-01-05|827.92|0"].amount == Decimal("827.92")

    comma = by_key["FC1|2026-01-06|1700.49|0"]
    assert comma.amount == Decimal("1700.49")
    assert comma.link == "geha.pdf"

    assert "117130664J|2026-03-09|90.00|0" in by_key
    assert "117130664J|2026-03-09|90.00|1" in by_key
    assert by_key["117130664J|2026-03-09|90.00|1"].link == "b.pdf"

    typo = next(row for row in parsed.rows if row.check_number == "1170019231")
    assert typo.check_date is None
    assert typo.deposit_date == date(2026, 3, 9)

    assert all(row.deposit_date.year == 2026 for row in parsed.rows)
    assert parsed.skipped_out_of_year == 1
    assert parsed.skipped_no_amount == 2
    assert parsed.duplicate_sheet_keys == 1
    assert len(parsed.errors) == 2
    assert any("103.003.64" in err.message for err in parsed.errors)
    assert any("Implausible check date" in err.message for err in parsed.errors)
    assert not any(row.payer == "Visa Payer" for row in parsed.rows)
    assert not any(row.payer == "Old" for row in parsed.rows)


def test_export_round_trip(tmp_path: Path):
    payload = export_checks_deposits_workbook(
        [
            {
                "check_date": "2026-01-02",
                "payer": "CSS LLC",
                "amount": "827.92",
                "check_number": "69017024",
                "deposit_date": "2026-01-05",
                "link": "BH-CSS.pdf",
                "notes": "ok",
            }
        ]
    )
    parsed = parse_checks_deposits_workbook(BytesIO(payload))
    assert len(parsed.rows) == 1
    row = parsed.rows[0]
    assert row.check_number == "69017024"
    assert row.amount == Decimal("827.92")
    assert row.deposit_date == date(2026, 1, 5)
    assert row.notes == "ok"
    book = load_workbook(BytesIO(payload))
    assert book.sheetnames == ["PT OF THE CITY"]
