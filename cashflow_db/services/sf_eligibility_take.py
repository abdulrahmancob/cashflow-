"""Copy Snowflake visit fields onto pending eligibility rows, then compare.

A row is updated only when the sheet's effective status is pending and
Snowflake has some other status. Rows that are not pending are left as they
are, and Snowflake visits that are not already on the sheet are not inserted.
"""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

from cashflow_db.repository import client
from cashflow_db.repository.eligibility import (
    OVERRIDE_FIELDS,
    SF_PAID_TOLERANCE,
    _as_number,
    _compare_value,
    iter_export_work_items,
    normalize_edit_value,
    normalize_visit_status,
    overlay_live_sf,
    parse_manual_overrides,
    sheet_export_headers,
    sheet_export_row,
)

YEAR_START = date(2026, 1, 1)
YEAR_END = date(2027, 1, 1)
EXPORT_PATH = Path("/data/exports/eligibility_vs_snowflake.xlsx")
TAKE_REASON = "copied from Snowflake because the sheet status was pending"

# Overlay keys -> sheet fields that manual overrides keep on the row.
_SF_SHEET_FIELDS: tuple[tuple[str, str], ...] = (
    ("sf_status", "source_visit_status"),
    ("sf_client_payment", "client_payment"),
    ("sf_insurance_payment", "insurance_payment"),
    ("sf_updated_payment", "updated_payment"),
    ("sf_coinsurance_payment", "coinsurance_payment"),
    ("sf_reduction", "reduction"),
    ("sf_adjusted", "adjusted"),
    ("sf_charged_amount", "charged_amount"),
    ("sf_details", "details"),
    ("sf_insurance_check_number", "insurance_check_number"),
    ("sf_insurance_check_date", "insurance_check_date"),
    ("sf_insurance_check_amount", "insurance_check_amount"),
    ("sf_secondary_check_number", "secondary_check_number"),
    ("sf_secondary_check_date", "secondary_check_date"),
    ("sf_secondary_check_amount", "secondary_check_amount"),
    ("sf_collector_1", "collector_1"),
    ("sf_posting_date_1", "posting_date_1"),
    ("sf_collector_2", "collector_2"),
    ("sf_posting_date_2", "posting_date_2"),
    ("sf_collector_3", "collector_3"),
    ("sf_posting_date_3", "posting_date_3"),
    ("sf_visit_status_sheet", "visit_status_sheet"),
    ("sf_corrected", "corrected"),
    ("sf_corrected_date", "corrected_date"),
    ("sf_visit_id", "sf_visit_id"),
    ("sf_insurance_id", "insurance_id"),
    ("sf_secondary_insurance", "secondary_insurance"),
    ("sf_secondary_insurance_id", "secondary_insurance_id"),
    ("sf_updated_check_number", "updated_check_number"),
    ("sf_updated_check_date", "updated_check_date"),
    ("sf_updated_check_amount", "updated_check_amount"),
    ("sf_fourth_check_number", "fourth_check_number"),
    ("sf_fourth_check_date", "fourth_check_date"),
    ("sf_fourth_check_amount", "fourth_check_amount"),
    ("sf_work_status", "work_status"),
    ("sf_work_date", "work_date"),
    ("sf_denial_reason", "denial_reason"),
    ("sf_root_cause", "root_cause"),
    ("sf_actions_taken", "actions_taken"),
    ("sf_collection_status", "collection_status"),
)

_DIFF_HEADERS = ("EMR", "DOS", "Difference", "Eligibility", "Snowflake")
_BATCH = 400


def should_copy_visit(local_status: Any, sf_status: Any) -> bool:
    """True when the sheet is pending and Snowflake has a different status."""
    if normalize_visit_status(local_status) != "pending":
        return False
    raw = str(sf_status or "").strip().lower()
    if not raw or raw == "blank":
        return False
    return normalize_visit_status(sf_status) != "pending"


def sheet_values_from_sf(fields: dict[str, Any]) -> dict[str, Any]:
    """Map Snowflake overlay fields onto sheet columns. Empty values are omitted."""
    out: dict[str, Any] = {}
    for sf_key, sheet_key in _SF_SHEET_FIELDS:
        raw = fields.get(sf_key)
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            continue
        if sheet_key == "source_visit_status":
            status = normalize_visit_status(raw)
            if status == "pending":
                continue
            out[sheet_key] = status
            continue
        if sheet_key not in OVERRIDE_FIELDS:
            continue
        try:
            value = normalize_edit_value(sheet_key, raw)
        except ValueError:
            continue
        if value is None:
            continue
        out[sheet_key] = value
    if "insurance_payment" in out:
        out["paid_amount"] = out["insurance_payment"]
    if "insurance_check_number" in out:
        out["check_number"] = out["insurance_check_number"]
    if "insurance_check_date" in out:
        out["check_date"] = out["insurance_check_date"]
    return out


def _money(value: Any) -> float:
    number = _as_number(value)
    if number is None:
        return 0.0
    return round(float(number), 2)


def remaining_differences(
    *,
    emr: str,
    dos: str,
    local_status: Any | None,
    local_amount: Any | None,
    sf_status: Any | None,
    sf_amount: Any | None,
    on_sheet: bool,
    on_snowflake: bool,
) -> list[tuple[str, str, str, str, str]]:
    """Differences that remain after the copy. A matching copied row returns nothing."""
    if on_sheet and not on_snowflake:
        return [(emr, dos, "only-on-sheet", _status_text(local_status), "")]
    if on_snowflake and not on_sheet:
        return [(emr, dos, "only-on-snowflake", "", _status_text(sf_status))]
    rows: list[tuple[str, str, str, str, str]] = []
    local_norm = normalize_visit_status(local_status)
    sf_norm = normalize_visit_status(sf_status)
    if local_norm != sf_norm:
        rows.append((emr, dos, "status", local_norm, sf_norm))
    if abs(_money(local_amount) - _money(sf_amount)) > SF_PAID_TOLERANCE:
        rows.append(
            (emr, dos, "amount", f"{_money(local_amount):.2f}", f"{_money(sf_amount):.2f}")
        )
    return rows


def _status_text(value: Any) -> str:
    text = str(value or "").strip()
    return normalize_visit_status(text) if text else ""


def _dos_iso(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    return None


def _chunks(rows: list[dict[str, Any]], size: int) -> Iterator[list[dict[str, Any]]]:
    for start in range(0, len(rows), size):
        yield rows[start : start + size]


def _load_pending_matches(conn) -> list[dict[str, Any]]:
    return client.fetchall(
        conn,
        """
        SELECT wi.work_item_id::text AS work_item_id,
               btrim(wi.emr_patient_id) AS emr_patient_id,
               wi.dos,
               wi.source_visit_status,
               wi.manual_overrides,
               sf.patient_name AS sf_patient_name,
               sf.insurance AS sf_insurance,
               sf.status AS sf_status
        FROM ops.eligibility_work_item wi
        JOIN analytics.snowflake_visit_kpi sf
          ON sf.emr_id = btrim(wi.emr_patient_id)
         AND sf.date_of_service = wi.dos
        WHERE lower(btrim(coalesce(
                nullif(btrim(wi.manual_overrides->>'source_visit_status'), ''),
                nullif(btrim(wi.source_visit_status), ''),
                'pending'
              ))) IN ('pending', 'blank')
          AND lower(btrim(coalesce(sf.status, ''))) NOT IN ('', 'pending', 'blank')
          AND btrim(wi.emr_patient_id) <> ''
        """,
    )


def _merge_overrides(current: Any, updates: dict[str, Any]) -> dict[str, Any]:
    merged = parse_manual_overrides(current)
    for key, value in updates.items():
        if key not in OVERRIDE_FIELDS or value is None:
            continue
        if _compare_value(key, merged.get(key)) == _compare_value(key, value) and key in merged:
            continue
        merged[key] = value
    return merged


def take_pending_from_snowflake(conn) -> dict[str, int]:
    """Persist Snowflake fields onto pending sheet rows. Does not insert visits."""
    matches = _load_pending_matches(conn)
    print(f"snowflake take candidates {len(matches)}", flush=True)
    updated = 0
    for batch in _chunks(matches, _BATCH):
        rows = [
            {
                "work_item_id": item["work_item_id"],
                "emr_patient_id": item["emr_patient_id"],
                "dos": item["dos"],
                "source_visit_status": item["source_visit_status"],
                "manual_overrides": item["manual_overrides"],
                "sf_patient_name": item.get("sf_patient_name"),
                "sf_insurance": item.get("sf_insurance"),
                "sf_status_sql": item.get("sf_status"),
            }
            for item in batch
            if should_copy_visit(
                (parse_manual_overrides(item.get("manual_overrides")).get("source_visit_status")
                 or item.get("source_visit_status")),
                item.get("sf_status"),
            )
        ]
        if not rows:
            continue
        overlay_live_sf(conn, rows)
        params: list[tuple[Any, ...]] = []
        history: list[tuple[Any, ...]] = []
        for row in rows:
            updates = sheet_values_from_sf(row)
            patient = str(row.get("sf_patient_name") or "").strip()
            insurance = str(row.get("sf_insurance") or "").strip()
            if patient:
                updates["patient_name"] = patient
            if insurance:
                updates["insurance_name"] = insurance
            status = updates.get("source_visit_status")
            if not status:
                continue
            merged = _merge_overrides(row.get("manual_overrides"), updates)
            old_status = normalize_visit_status(
                parse_manual_overrides(row.get("manual_overrides")).get("source_visit_status")
                or row.get("source_visit_status")
            )
            params.append(
                (
                    status,
                    patient or None,
                    insurance or None,
                    json.dumps(merged, default=str),
                    row["work_item_id"],
                )
            )
            history.append((row["work_item_id"], old_status, status, TAKE_REASON))
        if not params:
            continue
        client.executemany(
            conn,
            """
            UPDATE ops.eligibility_work_item
            SET source_visit_status = %s,
                patient_name = COALESCE(%s, patient_name),
                insurance_name = COALESCE(%s, insurance_name),
                manual_overrides = %s::jsonb,
                updated_at = now()
            WHERE work_item_id = %s::uuid
            """,
            params,
        )
        client.executemany(
            conn,
            """
            INSERT INTO ops.eligibility_history (
                work_item_id, column_name, old_value, new_value, reason_text
            )
            VALUES (%s::uuid, 'source_visit_status', %s, %s, %s)
            """,
            history,
        )
        updated += len(params)
        print(f"snowflake take updated {updated}", flush=True)
    return {"candidates": len(matches), "updated": updated}


def _load_sf_index(conn) -> dict[tuple[str, str], tuple[str, Any]]:
    rows = client.fetchall(
        conn,
        """
        SELECT btrim(emr_id) AS emr_id,
               date_of_service,
               status,
               insurance_payment
        FROM analytics.snowflake_visit_kpi
        WHERE date_of_service >= %s
          AND date_of_service < %s
          AND btrim(coalesce(emr_id, '')) <> ''
        """,
        (YEAR_START, YEAR_END),
    )
    index: dict[tuple[str, str], tuple[str, Any]] = {}
    for row in rows:
        dos = _dos_iso(row.get("date_of_service"))
        emr = str(row.get("emr_id") or "").strip()
        if not emr or not dos:
            continue
        index[(emr, dos)] = (row.get("status"), row.get("insurance_payment"))
    return index


def _iter_snowflake_sheet_rows(conn) -> Iterator[dict[str, Any]]:
    keys = client.fetchall(
        conn,
        """
        SELECT btrim(emr_id) AS emr_patient_id,
               date_of_service AS dos,
               patient_name,
               insurance AS insurance_name
        FROM analytics.snowflake_visit_kpi
        WHERE date_of_service >= %s
          AND date_of_service < %s
          AND btrim(coalesce(emr_id, '')) <> ''
        ORDER BY date_of_service, emr_id
        """,
        (YEAR_START, YEAR_END),
    )
    for batch in _chunks(keys, _BATCH):
        rows = [dict(item) for item in batch]
        overlay_live_sf(conn, rows)
        for row in rows:
            filled = sheet_values_from_sf(row)
            filled["emr_patient_id"] = row.get("emr_patient_id")
            filled["dos"] = row.get("dos")
            filled["patient_name"] = row.get("patient_name")
            filled["insurance_name"] = row.get("insurance_name")
            yield filled


def write_compare_workbook(conn, path: Path | None = None) -> dict[str, int]:
    """Write Eligibility, Snowflake, and remaining Differences after the copy.

    Amount differences compare the sheet Total Amount with Snowflake's
    insurance payment.
    """
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl required") from exc

    dest = path or EXPORT_PATH
    dest.parent.mkdir(parents=True, exist_ok=True)
    sf_index = _load_sf_index(conn)
    seen: set[tuple[str, str]] = set()
    differences: list[tuple[str, str, str, str, str]] = []
    eligibility_rows = 0

    wb = Workbook(write_only=True)
    eligibility = wb.create_sheet("Eligibility")
    eligibility.append(list(sheet_export_headers()))
    for row in iter_export_work_items(conn):
        eligibility.append(sheet_export_row(row))
        eligibility_rows += 1
        dos = _dos_iso(row.get("dos"))
        emr = str(row.get("emr_patient_id") or "").strip()
        if not emr or not dos or not dos.startswith("2026-"):
            continue
        key = (emr, dos)
        sf = sf_index.get(key)
        if sf is None:
            differences.extend(
                remaining_differences(
                    emr=emr,
                    dos=dos,
                    local_status=row.get("source_visit_status"),
                    local_amount=row.get("total_amount"),
                    sf_status=None,
                    sf_amount=None,
                    on_sheet=True,
                    on_snowflake=False,
                )
            )
            continue
        seen.add(key)
        differences.extend(
            remaining_differences(
                emr=emr,
                dos=dos,
                local_status=row.get("source_visit_status"),
                local_amount=row.get("total_amount"),
                sf_status=sf[0],
                sf_amount=sf[1],
                on_sheet=True,
                on_snowflake=True,
            )
        )
        if eligibility_rows % 5000 == 0:
            print(f"eligibility sheet rows {eligibility_rows}", flush=True)

    snowflake = wb.create_sheet("Snowflake")
    snowflake.append(list(sheet_export_headers()))
    snowflake_rows = 0
    for row in _iter_snowflake_sheet_rows(conn):
        snowflake.append(sheet_export_row(row))
        snowflake_rows += 1
        if snowflake_rows % 5000 == 0:
            print(f"snowflake sheet rows {snowflake_rows}", flush=True)

    for key, (sf_status, sf_amount) in sf_index.items():
        if key in seen:
            continue
        differences.extend(
            remaining_differences(
                emr=key[0],
                dos=key[1],
                local_status=None,
                local_amount=None,
                sf_status=sf_status,
                sf_amount=sf_amount,
                on_sheet=False,
                on_snowflake=True,
            )
        )

    diff_sheet = wb.create_sheet("Differences")
    diff_sheet.append(list(_DIFF_HEADERS))
    for item in differences:
        diff_sheet.append(list(item))
    wb.save(dest)
    print(f"wrote {dest}", flush=True)
    return {
        "eligibility_rows": eligibility_rows,
        "snowflake_rows": snowflake_rows,
        "difference_rows": len(differences),
    }


def run_take_and_workbook(path: Path | None = None, *, take: bool = True) -> dict[str, Any]:
    from cashflow_db.repository import connection

    dest = path or EXPORT_PATH
    with connection() as conn:
        taken = take_pending_from_snowflake(conn) if take else {"candidates": 0, "updated": 0}
    with connection() as conn:
        workbook = write_compare_workbook(conn, dest)
    return {"take": taken, "workbook": workbook, "path": str(dest)}


if __name__ == "__main__":
    import sys

    out = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    print(json.dumps(run_take_and_workbook(path=out), indent=2, default=str))
