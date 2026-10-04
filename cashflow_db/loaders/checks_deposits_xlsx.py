"""Checks & Deposits xlsx parse/export.

Deposit Date is the year gate (default 2026). Sheet titles are layout only.
Check numbers are not unique, so each row gets a stable sheet_key:
``{check_number}|{deposit_date}|{amount}|{seq}``.
"""

from __future__ import annotations

import io
import re
from dataclasses import asdict, dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, BinaryIO

from cashflow_db.loaders.load_checks_deposits import normalize_check_number
from cashflow_db.util import parse_date, safe_str

try:
    import openpyxl
    from openpyxl import Workbook
except ImportError:  # pragma: no cover
    openpyxl = None
    Workbook = None  # type: ignore[misc, assignment]

_WS_RE = re.compile(r"\s+")
_MONEY_ABS_MAX = Decimal("999999999999.99")
_CENT = Decimal("0.01")

ALLOWED_YEARS = frozenset({2026})

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "check_date": ("column 1", "check date", "date", "received date"),
    "payer": ("source(insurance)", "source", "insurance", "payer"),
    "amount": ("amount",),
    "check_number": ("check number", "check #", "check", "check/reference"),
    "deposit_date": ("deposit date",),
    "link": ("link",),
    "notes": ("notes",),
}

CANONICAL_HEADERS = [
    "Check Date",
    "Source(insurance)",
    "Amount",
    "Check number",
    "Deposit Date",
    "Link",
    "Notes",
]

SHEET_NAME = "PT OF THE CITY"


def normalize_header(cell: Any) -> str:
    if cell is None:
        return ""
    text = str(cell).replace("\xa0", " ").replace("\u200b", "")
    return _WS_RE.sub(" ", text).strip().lower()


def parse_sheet_money(value: Any) -> tuple[Decimal | None, str | None]:
    """Parse a sheet amount.

    A single comma with 1–2 decimal digits is a decimal separator (``1700,49``).
    More than one dot is ambiguous (``103.003.64``) and returns an error instead
    of a guessed number. Empty values return ``(None, None)``.
    """
    if value is None or value == "":
        return None, None
    if isinstance(value, bool):
        return None, None
    if isinstance(value, Decimal):
        amount = value
    elif isinstance(value, (int, float)):
        amount = Decimal(str(value))
    else:
        text = str(value).strip()
        if not text or text.upper() in {"#N/A", "N/A", "NA", "-", "NONE", "NULL"}:
            return None, None
        neg = text.startswith("(") and text.endswith(")")
        if neg:
            text = text[1:-1].strip()
        text = text.replace("$", "").replace(" ", "").replace("\xa0", "")
        if text.count(".") > 1:
            return None, f"Ambiguous amount {value!r}"
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            if re.fullmatch(r"-?\d{1,3}(?:,\d{3})+", text):
                text = text.replace(",", "")
            elif re.fullmatch(r"-?\d+,\d{1,2}", text):
                text = text.replace(",", ".")
            else:
                return None, f"Ambiguous amount {value!r}"
        if not re.fullmatch(r"-?\d+(?:\.\d+)?", text or ""):
            return None, f"Invalid amount {value!r}"
        try:
            amount = Decimal(text)
        except InvalidOperation:
            return None, f"Invalid amount {value!r}"
        if neg:
            amount = -amount
    if abs(amount) > _MONEY_ABS_MAX:
        return None, f"Amount out of range {value!r}"
    return amount.quantize(_CENT, rounding=ROUND_HALF_UP), None


def sheet_key_for(
    check_number: str | None,
    deposit_date: date,
    amount: Decimal,
    seq: int,
) -> str:
    number = check_number or ""
    return f"{number}|{deposit_date.isoformat()}|{amount:.2f}|{seq}"


def _build_col_map(header_row: tuple[Any, ...] | list[Any]) -> dict[str, int]:
    norm_to_idx: dict[str, int] = {}
    for i, cell in enumerate(header_row):
        key = normalize_header(cell)
        if not key or key in norm_to_idx:
            continue
        norm_to_idx[key] = i

    col_map: dict[str, int] = {}
    for field_name, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in norm_to_idx:
                col_map[field_name] = norm_to_idx[alias]
                break
    if "check_date" not in col_map and header_row:
        col_map["check_date"] = 0
    return col_map


def _cell(row: tuple[Any, ...] | list[Any], idx: int | None) -> Any:
    if idx is None or idx < 0 or idx >= len(row):
        return None
    return row[idx]


@dataclass
class ParsedCheckRow:
    sheet_key: str
    check_date: date | None
    payer: str | None
    amount: Decimal
    check_number: str | None
    deposit_date: date
    deposit_month: date
    link: str | None = None
    notes: str | None = None
    source_sheet: str | None = None
    source_row: int | None = None

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        for key, value in list(data.items()):
            if isinstance(value, date):
                data[key] = value.isoformat()
            elif isinstance(value, Decimal):
                data[key] = str(value)
        return data


@dataclass
class SheetParseError:
    sheet: str
    message: str
    row: int | None = None


@dataclass
class ParseResult:
    rows: list[ParsedCheckRow] = field(default_factory=list)
    errors: list[SheetParseError] = field(default_factory=list)
    skipped_sheets: list[str] = field(default_factory=list)
    skipped_out_of_year: int = 0
    skipped_no_amount: int = 0
    skipped_no_deposit_date: int = 0
    duplicate_sheet_keys: int = 0

    @property
    def month_bounds(self) -> list[tuple[date, date]]:
        import calendar

        months: set[date] = set()
        for row in self.rows:
            months.add(row.deposit_date.replace(day=1))
        bounds: list[tuple[date, date]] = []
        for start in sorted(months):
            last = calendar.monthrange(start.year, start.month)[1]
            bounds.append((start, date(start.year, start.month, last)))
        return bounds


def _implausible_check_date(value: date | None) -> bool:
    return value is not None and (value.year < 1990 or value.year > 2100)


def parse_checks_deposits_workbook(
    source: Path | str | BinaryIO | bytes,
    *,
    years: set[int] | frozenset[int] | None = None,
) -> ParseResult:
    """Parse Checks & Deposits sheets. Rows outside ``years`` are counted and dropped."""
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to parse Checks and Deposits")

    allowed = frozenset(years) if years is not None else ALLOWED_YEARS
    result = ParseResult()
    if isinstance(source, (bytes, bytearray)):
        wb = openpyxl.load_workbook(io.BytesIO(source), read_only=True, data_only=True)
    elif hasattr(source, "read"):
        wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    else:
        wb = openpyxl.load_workbook(Path(source), read_only=True, data_only=True)

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            header = next(rows_iter, None)
            if header is None:
                result.skipped_sheets.append(sheet_name)
                continue
            col_map = _build_col_map(header)
            if "deposit_date" not in col_map or "amount" not in col_map:
                result.skipped_sheets.append(sheet_name)
                continue

            seen: dict[str, int] = {}
            for row_num, raw in enumerate(rows_iter, start=2):
                if not raw or all(c is None or str(c).strip() == "" for c in raw):
                    continue
                deposit_date = parse_date(_cell(raw, col_map.get("deposit_date")))
                if deposit_date is None:
                    result.skipped_no_deposit_date += 1
                    continue
                if deposit_date.year not in allowed:
                    result.skipped_out_of_year += 1
                    continue

                amount, amount_error = parse_sheet_money(_cell(raw, col_map.get("amount")))
                check_number = normalize_check_number(
                    _cell(raw, col_map.get("check_number"))
                )
                if amount_error:
                    result.errors.append(
                        SheetParseError(
                            sheet=sheet_name,
                            row=row_num,
                            message=f"{amount_error} for check {check_number or '—'}",
                        )
                    )
                    continue
                if amount is None:
                    result.skipped_no_amount += 1
                    continue

                check_date = parse_date(_cell(raw, col_map.get("check_date")))
                if _implausible_check_date(check_date):
                    result.errors.append(
                        SheetParseError(
                            sheet=sheet_name,
                            row=row_num,
                            message=(
                                f"Implausible check date {check_date.isoformat()} "
                                f"for check {check_number or '—'}; stored blank"
                            ),
                        )
                    )
                    check_date = None

                base = f"{check_number or ''}|{deposit_date.isoformat()}|{amount:.2f}"
                seq = seen.get(base, 0)
                seen[base] = seq + 1
                if seq:
                    result.duplicate_sheet_keys += 1

                result.rows.append(
                    ParsedCheckRow(
                        sheet_key=sheet_key_for(check_number, deposit_date, amount, seq),
                        check_date=check_date,
                        payer=safe_str(_cell(raw, col_map.get("payer"))),
                        amount=amount,
                        check_number=check_number,
                        deposit_date=deposit_date,
                        deposit_month=deposit_date.replace(day=1),
                        link=safe_str(_cell(raw, col_map.get("link"))),
                        notes=safe_str(_cell(raw, col_map.get("notes"))),
                        source_sheet=sheet_name,
                        source_row=row_num,
                    )
                )
    finally:
        wb.close()
    return result


def export_checks_deposits_workbook(rows: list[dict[str, Any]]) -> bytes:
    """Build one PT OF THE CITY sheet in the source column order."""
    if Workbook is None:
        raise RuntimeError("openpyxl is required to export Checks and Deposits")

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(CANONICAL_HEADERS)

    def sort_key(row: dict[str, Any]) -> tuple[str, str, str]:
        return (
            str(row.get("deposit_date") or ""),
            str(row.get("check_number") or ""),
            str(row.get("sheet_key") or ""),
        )

    for row in sorted(rows, key=sort_key):
        amount = row.get("amount")
        ws.append(
            [
                row.get("check_date"),
                row.get("payer"),
                float(amount) if amount is not None and amount != "" else None,
                row.get("check_number"),
                row.get("deposit_date"),
                row.get("link"),
                row.get("notes"),
            ]
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
