"""Build intake_census_stats.xlsx from a census run directory.

Inputs (all under --dir, default /data/exports/intake_census_v2): patients.csv, summary.json,
rows.jsonl; optional patients_named.csv.gz (webpt_patient_id, patient_name) for names.
Output: --out (default <dir>/intake_census_stats.xlsx).
"""

from __future__ import annotations

import argparse
import csv
import gzip
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from intake_reader import GROUPS  # noqa: E402

DESCRIPTIONS = {
    "doctor": "Doctor referral / recommendation",
    "google": "Google",
    "website": "Website",
    "zocdoc": "Zocdoc",
    "social_media": "Social media",
    "insurance": "Insurance recommendation",
    "friend_family": "Friend / family / word of mouth",
    "event": "Event or community outreach",
    "flyer_doctor_office": "Flyer from a doctor's office",
    "flyer_street": "Flyer from street distribution",
    "marketing_table": "Marketing table",
    "direct_mail": "Direct mail",
    "clinic_staff": "Clinic staff",
    "walk_in": "Walk-in / passed by",
    "lives_nearby": "Lives nearby",
    "phone": "Phone",
    "other": "Other (text kept in Other text)",
    "multiple": "More than one option marked (see Marks)",
    "unmarked": "Question found, nothing marked",
    "no_question": "Form has no referral question",
    "unreadable": "Could not be read (see Reason)",
}
FAMILY_NAMES = {
    "old_checkbox": "Checkbox form (Please check what applies)",
    "new_circle": "Circle form (How did you find us)",
    "es_checkbox": "Spanish checkbox form",
    "es_circle": "Spanish circle form",
    "old_bullet": "Bullet form (Welcome to PT of The City)",
    "tiny": "One-line form",
    "generic": "Other layout",
    "": "No question found",
    "none": "No question found",
}
LANG_NAMES = {"english": "English", "spanish": "Spanish", "bilingual": "Bilingual", "unknown": "Unknown"}
METHOD_NAMES = {"ocr": "OCR (scanned image)", "native_text": "Native PDF text", "mixed": "Mixed", "empty": "Empty", "error": "Error"}


def load_names(path: Path | None) -> dict[str, str]:
    names: dict[str, str] = {}
    if not path or not path.is_file():
        return names
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            pid = (row.get("webpt_patient_id") or "").strip()
            name = (row.get("patient_name") or "").strip()
            if pid and name:
                names[pid] = name
    return names


def names_from_db() -> dict[str, str]:
    """Current patient names from core.patient / core.patient_history (read-only)."""
    import os

    url = os.environ.get("CASHFLOW_DATABASE_URL", "").strip() or os.environ.get("DATABASE_URL", "").strip()
    names: dict[str, str] = {}
    if not url:
        return names
    try:
        import psycopg

        with psycopg.connect(url) as conn:
            conn.read_only = True
            with conn.cursor() as cur:
                cur.execute(
                    """
                    SELECT DISTINCT ON (p.webpt_patient_id)
                           p.webpt_patient_id,
                           COALESCE(NULLIF(btrim(ph.patient_name), ''), p.name_key) AS patient_name
                    FROM core.patient p
                    LEFT JOIN core.patient_history ph
                      ON ph.patient_id = p.patient_id AND ph.is_current
                    WHERE p.webpt_patient_id IS NOT NULL
                      AND btrim(p.webpt_patient_id) <> ''
                    ORDER BY p.webpt_patient_id, ph.valid_from DESC NULLS LAST
                    """
                )
                for webpt_id, name in cur.fetchall():
                    key = str(webpt_id or "").strip()
                    if key and name and key not in names:
                        names[key] = str(name).strip()
    except Exception as exc:  # names are a convenience, never a reason to fail the workbook
        print(f"names_from_db skipped: {type(exc).__name__}", flush=True)
    return names


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dir", default="/data/exports/intake_census_v2")
    parser.add_argument("--names", default="")
    parser.add_argument("--out", default="")
    args = parser.parse_args()
    base = Path(args.dir)
    out = Path(args.out) if args.out else base / "intake_census_stats.xlsx"
    names_path = Path(args.names) if args.names else (base / "patients_named.csv.gz")
    names = load_names(names_path if names_path.is_file() else None)
    if not names:
        names = names_from_db()
    print(f"names={len(names)}", flush=True)
    summary = json.loads((base / "summary.json").read_text(encoding="utf-8"))
    with (base / "patients.csv").open(encoding="utf-8", newline="") as handle:
        patients = list(csv.DictReader(handle))
    files_by_patient: dict[str, list[dict]] = defaultdict(list)
    rows_path = base / "rows.jsonl"
    file_rows = []
    if rows_path.is_file():
        for line in rows_path.read_text(encoding="utf-8", errors="replace").splitlines():
            try:
                file_rows.append(json.loads(line))
            except json.JSONDecodeError:
                continue

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")

    def sheet(title: str, header: list[str], rows: list[list], widths: dict[int, int] | None = None):
        ws = wb.create_sheet(title) if wb.worksheets and wb.worksheets[0].title != "Sheet" else wb.active
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for index, name in enumerate(header, 1):
            width = (widths or {}).get(index, max(12, min(48, len(str(name)) + 4)))
            ws.column_dimensions[get_column_letter(index)].width = width
        return ws

    with_intake = [p for p in patients if p.get("has_intake") == "yes"]
    total = len(patients)
    n_with = len(with_intake)
    group_counts = Counter(p.get("source_group") or "Unreadable" for p in with_intake)
    code_counts = Counter(p.get("source") or "unreadable" for p in with_intake)
    known = sum(v for k, v in code_counts.items() if k not in ("unmarked", "no_question", "unreadable", "multiple"))
    review = sum(1 for p in with_intake if p.get("needs_review"))
    files_total = summary.get("intake_files", len(file_rows))

    def share(n: int, d: int) -> float | None:
        return round(n / d, 4) if d else None

    summary_rows = [
        ["Total patients", total, None],
        ["Patients with an intake", n_with, share(n_with, total)],
        ["Patients without an intake", total - n_with, share(total - n_with, total)],
        ["Intake files read", files_total, None],
        ["Intake files not linked to a patient", summary.get("unmatched_intake_files", 0), None],
        ["Patients with a known referral source", known, share(known, n_with)],
        ["Patients with more than one mark (see Marks)", code_counts.get("multiple", 0), share(code_counts.get("multiple", 0), n_with)],
        ["Patients: question found, nothing marked", code_counts.get("unmarked", 0), share(code_counts.get("unmarked", 0), n_with)],
        ["Patients: form has no referral question", code_counts.get("no_question", 0), share(code_counts.get("no_question", 0), n_with)],
        ["Patients: unreadable (see Review sheet)", code_counts.get("unreadable", 0), share(code_counts.get("unreadable", 0), n_with)],
        ["Patients flagged for review", review, share(review, n_with)],
        ["Form layouts", len(summary.get("family", {})), None],
        [None, None, None],
        ["Counts are per patient. A patient with several intakes gets the reading with the best evidence; files that disagree are reported as 'Multiple marks' and flagged for review. Booking channel (phone, Zocdoc, website, walk-in) is a separate column and never counts as a referral source.", None, None],
    ]
    sheet("Summary", ["Item", "Count", "Share"], summary_rows, {1: 70, 2: 12, 3: 10})

    patient_rows = []
    for p in patients:
        pid = p["webpt_patient_id"]
        has = "Yes" if p.get("has_intake") == "yes" else "No"
        code = p.get("source") or ""
        patient_rows.append(
            [
                pid,
                names.get(pid, ""),
                has,
                int(p.get("file_count") or 0),
                (GROUPS.get(code, code) if has == "Yes" else "No intake"),
                code if has == "Yes" else "",
                p.get("marks") or "",
                p.get("other_text") or "",
                p.get("booking") or "",
                float(p.get("confidence")) if p.get("confidence") else None,
                "Yes" if p.get("needs_review") else ("" if has == "No" else "No"),
                LANG_NAMES.get(p.get("language") or "", p.get("language") or ""),
                FAMILY_NAMES.get(p.get("family") or "", p.get("family") or ""),
                p.get("shape_id") or "",
                METHOD_NAMES.get(p.get("text_method") or "", p.get("text_method") or ""),
                p.get("reason") or "",
            ]
        )
    sheet(
        "Patients",
        ["WebPT Patient ID", "Patient Name", "Has Intake", "Intake Files", "Referral Source", "Source Code", "Marks", "Other text", "Booking", "Confidence", "Needs Review", "Language", "Form Family", "Form Shape ID", "Read Method", "Reason"],
        patient_rows,
        {1: 16, 2: 28, 5: 24, 6: 16, 7: 22, 8: 36, 9: 14, 13: 34, 16: 30},
    )

    group_rows = [[g, n, share(n, n_with)] for g, n in sorted(group_counts.items(), key=lambda kv: -kv[1])]
    group_rows.append(["Total", n_with, None])
    sheet("Referral Source", ["Group", "Patients", "Share"], group_rows, {1: 30})
    code_rows = [[c, DESCRIPTIONS.get(c, c), GROUPS.get(c, c), n, share(n, n_with)] for c, n in sorted(code_counts.items(), key=lambda kv: -kv[1])]
    code_rows.append(["Total", "", "", n_with, None])
    sheet("Source Detail", ["Code", "Description", "Group", "Patients", "Share"], code_rows, {1: 20, 2: 44, 3: 24})

    lang_counts = Counter(p.get("language") or "unknown" for p in with_intake)
    lang_rows = [[k, LANG_NAMES.get(k, k), n, share(n, n_with)] for k, n in sorted(lang_counts.items(), key=lambda kv: -kv[1])]
    lang_rows.append(["", "Total", n_with, None])
    sheet("Language", ["Code", "Description", "Patients", "Share"], lang_rows)

    by_lang: dict[str, Counter[str]] = defaultdict(Counter)
    for p in with_intake:
        by_lang[p.get("source_group") or "Unreadable"][p.get("language") or "unknown"] += 1
    langs = ["english", "spanish", "bilingual", "unknown"]
    sbl_rows = [[g] + [by_lang[g].get(l, 0) for l in langs] + [sum(by_lang[g].values())] for g in sorted(by_lang, key=lambda g: -sum(by_lang[g].values()))]
    sbl_rows.append(["Total"] + [sum(by_lang[g].get(l, 0) for g in by_lang) for l in langs] + [n_with])
    sheet("Source by Language", ["Referral source", "English", "Spanish", "Bilingual", "Unknown", "Total"], sbl_rows, {1: 30})

    fam_counts = Counter(p.get("family") or "none" for p in with_intake)
    fam_rows = [[FAMILY_NAMES.get(k, k), k, n, share(n, n_with)] for k, n in sorted(fam_counts.items(), key=lambda kv: -kv[1])]
    sheet("Form Families", ["Form family", "Code", "Patients", "Share"], fam_rows, {1: 44})

    fc = Counter(int(p.get("file_count") or 0) for p in with_intake)
    sheet("Intakes per Patient", ["Intake files", "Patients"], [[k, n] for k, n in sorted(fc.items())])

    method_counts = Counter(p.get("text_method") or "" for p in with_intake)
    sheet("Read Method", ["Code", "Description", "Patients", "Share"], [[k, METHOD_NAMES.get(k, k), n, share(n, n_with)] for k, n in method_counts.most_common()])

    review_rows = []
    for p in with_intake:
        if not p.get("needs_review"):
            continue
        pid = p["webpt_patient_id"]
        review_rows.append([pid, names.get(pid, ""), GROUPS.get(p.get("source") or "", p.get("source") or ""), p.get("marks") or "", p.get("other_text") or "", float(p.get("confidence")) if p.get("confidence") else None, p.get("reason") or "", FAMILY_NAMES.get(p.get("family") or "", p.get("family") or "")])
    review_rows.sort(key=lambda r: (r[2], r[0]))
    sheet("Review", ["WebPT Patient ID", "Patient Name", "Reading", "Marks", "Other text", "Confidence", "Why it needs a look", "Form family"], review_rows, {2: 28, 3: 22, 4: 22, 5: 36, 7: 40, 8: 34})

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 110
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"WORKBOOK {out} patients={total} with_intake={n_with} review={review} groups={dict(group_counts)}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
